import tkinter as tk
from tkinter import messagebox
from openpyxl import Workbook, load_workbook
import os

# Create main window
root = tk.Tk()
root.title("Dark Theme Registration Form")
root.geometry("400x400")
root.config(bg="#1e1e1e")

# Fonts and colors
FONT = ("Segoe UI", 10)
LABEL_COLOR = "#ffffff"
ENTRY_BG = "#2d2d2d"
ENTRY_FG = "#ffffff"
BUTTON_BG = "#3a3a3a"
BUTTON_FG = "#00ffcc"

# Path to save Excel file - change username if needed
file_path = r"C:\Users\91912\Documents\users.xlsx"
folder = os.path.dirname(file_path)

# Function to register and save to Excel
def register():
    name = name_entry.get()
    email = email_entry.get()
    password = pass_entry.get()
    confirm = confirm_entry.get()

    if not name or not email or not password or not confirm:
        messagebox.showerror("Error", "All fields are required!")
        return

    if password != confirm:
        messagebox.showerror("Error", "Passwords do not match!")
        return

    # Create folder if it doesn't exist
    if not os.path.exists(folder):
        os.makedirs(folder)

    try:
        if os.path.exists(file_path):
            workbook = load_workbook(file_path)
            sheet = workbook.active
        else:
            workbook = Workbook()
            sheet = workbook.active
            # Add header row
            sheet.append(["Name", "Email", "Password"])

        # Append user data
        sheet.append([name, email, password])
        workbook.save(file_path)

        messagebox.showinfo("Success", f"Registered Successfully!\nWelcome, {name}")
        clear_form()

    except Exception as e:
        messagebox.showerror("Error", f"Failed to save data:\n{e}")

# Clear form entries
def clear_form():
    name_entry.delete(0, tk.END)
    email_entry.delete(0, tk.END)
    pass_entry.delete(0, tk.END)
    confirm_entry.delete(0, tk.END)

# UI Elements
tk.Label(root, text="Registration Form", font=("Segoe UI", 14, "bold"),
         fg=BUTTON_FG, bg="#1e1e1e").pack(pady=10)

form_frame = tk.Frame(root, bg="#1e1e1e")
form_frame.pack(pady=10)

# Field creation function
def add_field(label_text, is_password=False):
    tk.Label(form_frame, text=label_text, fg=LABEL_COLOR, bg="#1e1e1e", font=FONT).pack(anchor="w", pady=(5, 0))
    entry = tk.Entry(form_frame, font=FONT, fg=ENTRY_FG, bg=ENTRY_BG, insertbackground=ENTRY_FG,
                     show="*" if is_password else "")
    entry.pack(fill="x", pady=2)
    return entry

# Fields
name_entry = add_field("Full Name")
email_entry = add_field("Email")
pass_entry = add_field("Password", is_password=True)
confirm_entry = add_field("Confirm Password", is_password=True)

# Register button
tk.Button(root, text="Register", command=register, font=FONT,
          bg=BUTTON_BG, fg=BUTTON_FG, activebackground="#444", activeforeground=BUTTON_FG).pack(pady=20)

root.mainloop()
